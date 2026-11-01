import json
import zipfile

import pytest

from decision_forge.document_formats import check_office_archive
from decision_forge.sources import ingest_manifest, locate


def ingest(tmp_path, name, **options):
    manifest = tmp_path / "corpus.json"
    manifest.write_text(json.dumps({"sources": [{"id": "policy", "version": "v1", "path": name, **options}]}))
    return ingest_manifest(manifest)[0]


def test_docx_preserves_body_table_and_header_locations(tmp_path):
    docx = pytest.importorskip("docx")
    document = docx.Document()
    document.add_paragraph("Managers approve refunds above 100.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Role"
    table.cell(0, 1).text = "Limit"
    table.cell(1, 0).text = "Agent"
    table.cell(1, 1).text = "100"
    document.sections[0].header.paragraphs[0].text = "Internal policy"
    document.save(tmp_path / "policy.docx")
    result = ingest(tmp_path, "policy.docx", access={"classification": "restricted"})
    assert result.text.index("Managers") < result.text.index("Role") < result.text.index("Agent")
    assert any("table:2;row:2;column:2" in span.locator for span in result.spans)
    assert any("header" in span.locator for span in result.spans)
    assert result.access.classification == "restricted"
    assert locate(result, 0, 8) == "body;paragraph:1"


def test_xlsx_preserves_cells_and_formulas_without_evaluation(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Approval"
    sheet.append(["Role", "Limit"])
    sheet.append(["Manager", 1000])
    sheet["C2"] = "=B2*2"
    workbook.save(tmp_path / "policy.xlsx")
    result = ingest(tmp_path, "policy.xlsx")
    assert "B2: 1000" in result.text
    assert "[formula, not evaluated] =B2*2" in result.text
    assert result.spans[-1].locator == "sheet:Approval;cell:C2"


def test_pptx_text_and_table_keep_slide_coordinates(tmp_path):
    pptx = pytest.importorskip("pptx")
    from pptx.util import Inches

    presentation = pptx.Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Approve only after review."
    table = slide.shapes.add_table(1, 2, Inches(1), Inches(3), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Manager"
    presentation.save(tmp_path / "policy.pptx")
    result = ingest(tmp_path, "policy.pptx")
    assert "Approve only after review." in result.text
    assert result.spans[0].locator == "slide:1;shape:1;paragraph:1"
    assert result.spans[-1].locator == "slide:1;shape:2;row:1;column:2"


def test_html_extracts_static_content_without_scripts_or_hidden_nodes(tmp_path):
    pytest.importorskip("bs4")
    (tmp_path / "policy.html").write_text(
        "<html><head><title>Excluded</title></head><body><h1>Policy</h1><p>Review &amp; approve.</p>"
        '<script>fetch("https://invalid.example")</script><style>hidden</style>'
        "<p hidden>Secret</p><!-- Comment --><table><tr><td>Limit</td><td>100</td></tr></table></body></html>"
    )
    result = ingest(tmp_path, "policy.html")
    assert result.text == "Policy\nReview & approve.\nLimit\n100\n"
    assert all(span.locator.startswith("html:text-node:") for span in result.spans)


@pytest.mark.parametrize(
    "suffix,content",
    [("csv", 'Role,Rule\nManager,"Approve, then archive"\n'), ("tsv", "Role\tRule\nManager\tApprove\n")],
)
def test_delimited_formats_preserve_cell_locations(tmp_path, suffix, content):
    (tmp_path / f"policy.{suffix}").write_text(content)
    result = ingest(tmp_path, f"policy.{suffix}")
    assert "Manager\nApprove" in result.text
    assert result.spans[-1].locator == "row:2;column:2"


def test_auto_detects_workflow_and_rejects_legacy_binary(tmp_path):
    (tmp_path / "policy.json").write_text('{"rule":"Only managers approve."}')
    assert ingest(tmp_path, "policy.json").spans[0].locator == "json-pointer:/rule"
    (tmp_path / "policy.doc").write_bytes(b"legacy")
    with pytest.raises(ValueError, match="legacy"):
        ingest(tmp_path, "policy.doc")


def test_office_archive_expansion_limit(tmp_path):
    path = tmp_path / "large.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("large.xml", b"0" * 50_000_001)
    with pytest.raises(ValueError, match="expanded size"):
        check_office_archive(path)


def test_spreadsheet_sparse_dimensions_are_bounded(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    workbook.active["XFD1048576"] = "huge sparse sheet"
    workbook.save(tmp_path / "policy.xlsx")
    with pytest.raises(ValueError, match="traversal limit"):
        ingest(tmp_path, "policy.xlsx")


@pytest.mark.parametrize("sparse", [False, True])
def test_spreadsheet_incorrect_dimensions_do_not_hide_cells(tmp_path, sparse):
    openpyxl = pytest.importorskip("openpyxl")
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Handbook"
    workbook.active["A200002" if sparse else "A2"] = "Only managers approve refunds."
    original = tmp_path / "original.xlsx"
    workbook.save(original)
    with zipfile.ZipFile(original) as source, zipfile.ZipFile(tmp_path / "policy.xlsx", "w") as target:
        import re

        for item in source.infolist():
            data = source.read(item)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', data)
            target.writestr(item, data)
    if sparse:
        with pytest.raises(ValueError, match="traversal limit"):
            ingest(tmp_path, "policy.xlsx")
    else:
        result = ingest(tmp_path, "policy.xlsx")
        assert "Only managers approve refunds." in result.text
        assert result.spans[-1].locator.endswith("cell:A2")
