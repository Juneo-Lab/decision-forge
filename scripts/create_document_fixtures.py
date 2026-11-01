"""Create fictional native office files and image/PDF inputs for ingestion."""

from __future__ import annotations

import argparse
from pathlib import Path

from decision_forge.io import write_json


def create(output: Path) -> None:
    import docx
    import openpyxl
    import pptx
    from PIL import Image, ImageDraw, ImageFont
    from pptx.util import Inches

    if output.exists():
        raise ValueError("Choose a new fixture output directory")
    output.mkdir(parents=True)
    rule = "Only managers approve refunds above 100 dollars."
    document = docx.Document()
    document.add_heading("Fictional Northstar policy", 0)
    document.add_paragraph(rule)
    document.save(output / "policy.docx")
    workbook = openpyxl.Workbook()
    workbook.active.title = "Roles"
    workbook.active.append(["Role", "Permitted action"])
    workbook.active.append(["Manager", "Approve refunds above 100 dollars"])
    workbook.save(output / "roles.xlsx")
    slides = pptx.Presentation()
    slide = slides.slides.add_slide(slides.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = rule
    slides.save(output / "workflow.pptx")
    (output / "policy.html").write_text(
        f"<html><body><h1>Refund policy</h1><p>{rule}</p></body></html>", encoding="utf-8"
    )
    (output / "roles.csv").write_text(
        "Role,Permitted action\nManager,Approve refunds above 100 dollars\n", encoding="utf-8"
    )
    with Image.new("RGB", (1400, 180), "white") as image:
        ImageDraw.Draw(image).text((30, 40), rule, font=ImageFont.load_default(size=32), fill="black")
        image.save(output / "policy.png")
        image.save(output / "scanned-policy.pdf")
    sources = []
    for index, name in enumerate(
        ["policy.docx", "roles.xlsx", "workflow.pptx", "policy.html", "roles.csv", "policy.png", "scanned-policy.pdf"]
    ):
        source: dict[str, object] = {"id": f"source-{index}", "path": name, "version": "1", "adapter": "auto"}
        if name.endswith((".png", ".pdf")):
            source["ocr"] = {"backend": "rapidocr"}
        sources.append(source)
    # JSON is also valid YAML and is accepted by the manifest reader.
    write_json(output / "corpus.yaml", {"schema_version": "1.0", "sources": sources})


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    create(parser.parse_args().output)
